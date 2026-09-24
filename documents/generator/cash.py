import datetime
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.formatting.rule import CellIsRule
OUT="Haverton Recruitment/7 Launch Kit/Temporary Staffing Cash Flow.xlsx"
NAVY='1B2B45';GOLD='C9A55C';LIGHT='F3EFE6';INP='FFF6D9'
A=lambda **k:Font(name='Arial',size=11,**k)
thin=Side(style='thin',color='D9D2C3');bd=Border(top=thin,bottom=thin,left=thin,right=thin)
wb=Workbook();s=wb.active;s.title='Inputs'
s['A1']='Temporary Staffing Cash Flow (13 weeks)';s['A1'].font=Font(name='Arial',size=16,bold=True,color=NAVY)
s['A2']='Haverton Care Limited | Haverton Recruitment And Staffing | Planning estimate for discussion with your accountant';s['A2'].font=A(color='666666')
s['A3']='Change the yellow cells only. All amounts exclude VAT. Figures are estimates, not verified results.';s['A3'].font=A(bold=True,color='8A6D2F')
inputs=[
('First week of supply (Monday)',datetime.date(2027,4,5),'DD/MM/YYYY','Earliest realistic pilot date after the Go Live Gate. Change to your date.'),
('Opening cash available for temporary staffing (£)',10000,'£#,##0','[confirm] Cash you can set aside. Keep it separate from other business cash.'),
('Worker pay rate per hour (£)',14.50,'£#,##0.00','Must be at least the National Living Wage (£12.71 for 21+ from April 2026). Check local market rates.'),
('Charge rate to client per hour (£, ex VAT)',26.00,'£#,##0.00','Weekday day rate. Nights, weekends and bank holidays should be charged higher.'),
('Hours per worker per week',30,'0','Average hours actually worked and billed.'),
('Holiday pay accrual',0.1207,'0.00%','12.07% of pay is the usual accrual for irregular-hours workers.'),
('Employer National Insurance (planning load)',0.105,'0.0%','Blended planning figure, not the statutory calculation. Your payroll will calculate the exact amount.'),
('Employer pension (planning load)',0.025,'0.0%','Auto-enrolment minimum employer contribution is 3% of qualifying earnings.'),
('Other cost per hour worked (£)',0.75,'£#,##0.00','DBS, training, PPE, payroll software, per hour worked.'),
('Fixed weekly overheads for temporary staffing (£)',400,'£#,##0','[confirm] Insurance, on-call phone, software, compliance checks. Replace with real quotes.'),
('Client payment terms (days)',7,'0','Invoices go out weekly. Your plan is 7-day terms.'),
('Extra delay if clients pay late (weeks)',0,'0','Stress test: set to 2 or 4 to see the effect of late payers.'),
]
s['A5']='Input';s['B5']='Value';s['C5']='Notes'
for c in ('A5','B5','C5'): s[c].font=Font(name='Arial',size=11,bold=True,color='FFFFFF');s[c].fill=PatternFill('solid',fgColor=NAVY)
for i,(l,v,f,n) in enumerate(inputs,6):
    s.cell(row=i,column=1,value=l).font=A();c=s.cell(row=i,column=2,value=v);c.number_format=f;c.font=A(color='1F3F99');c.fill=PatternFill('solid',fgColor=INP);c.border=bd
    s.cell(row=i,column=3,value=n).font=A(color='555555');s.cell(row=i,column=3).alignment=Alignment(wrap_text=True)
# names: B6 date B7 cash B8 pay B9 charge B10 hours B11 hol B12 ni B13 pen B14 other B15 fixed B16 terms B17 delay
r=19
s.cell(row=r,column=1,value='Margin check per hour').font=Font(name='Arial',size=13,bold=True,color=NAVY)
calc=[('Loaded cost per hour (£)','=B8*(1+B11+B12+B13)+B14','£#,##0.00'),
('Gross profit per hour (£)','=B9-B20','£#,##0.00'),
('Gross margin','=IF(B9=0,0,B21/B9)','0.0%'),
('Meets Haverton floor (£5.50 per hour and 20%)?','=IF(AND(B21>=5.5,B22>=0.2),"Yes","No: needs Director approval")',None),
('Weeks between work and payment','=1+ROUNDUP(B16/7,0)+B17','0')]
for i,(l,f,nf) in enumerate(calc,20):
    s.cell(row=i,column=1,value=l).font=A();c=s.cell(row=i,column=2,value=f);c.font=A(bold=True);c.border=bd
    if nf:c.number_format=nf
s.cell(row=26,column=1,value='Results').font=Font(name='Arial',size=13,bold=True,color=NAVY)
res=[('Lowest cash balance in 13 weeks (£)',"=MIN('Cash Flow'!L3:L15)",'£#,##0'),
('Extra funding needed to stay above £0 (£)',"=MAX(0,-B27)",'£#,##0'),
('Week with lowest balance',"=INDEX('Cash Flow'!A3:A15,MATCH(B27,'Cash Flow'!L3:L15,0))",'0'),
('Unpaid client invoices at end of week 13 (£)',"=SUM('Cash Flow'!E3:E15)-SUM('Cash Flow'!F3:F15)",'£#,##0'),
('Gross profit earned over 13 weeks (£)',"=SUM('Cash Flow'!E3:E15)-SUM('Cash Flow'!G3:G15)-SUM('Cash Flow'!H3:H15)",'£#,##0')]
for i,(l,f,nf) in enumerate(res,27):
    s.cell(row=i,column=1,value=l).font=A();c=s.cell(row=i,column=2,value=f);c.font=A(bold=True);c.number_format=nf;c.border=bd
s.conditional_formatting.add('B28',CellIsRule(operator='greaterThan',formula=['0'],fill=PatternFill('solid',fgColor='F8D7D3'),font=Font(name='Arial',bold=True,color='9C1C1C')))
notes=['How to read this',
'You pay workers every week, one week in arrears. Clients pay you later. The gap is the cash you need.',
'Wages, holiday pay, National Insurance and pension are treated as paid with each weekly payroll. HMRC and pension payments are usually monthly, so this is slightly cautious.',
'VAT is left out. If you are VAT registered, the VAT you collect belongs to HMRC and must not be used to pay wages.',
'Ask your accountant to check this before you commit to any client. If the "Extra funding needed" cell is red, arrange funding (for example invoice finance) before the pilot.']
for i,t in enumerate(notes,33):
    c=s.cell(row=i,column=1,value=t);c.font=A(bold=(i==33),color=NAVY if i==33 else '333333')
s.column_dimensions['A'].width=52;s.column_dimensions['B'].width=18;s.column_dimensions['C'].width=90
c=wb.create_sheet('Cash Flow')
c['A1']='Weekly cash flow (change workers in the yellow column)';c['A1'].font=Font(name='Arial',size=14,bold=True,color=NAVY)
H=['Week','Week starting','Workers on shift','Hours billed','Invoiced to clients (£)','Cash in from clients (£)','Wage cost incurred (£)','Other costs (£)','Wages paid this week (£)','Fixed overheads (£)','Cash out total (£)','Closing cash (£)']
for j,h in enumerate(H,1):
    x=c.cell(row=2,column=j,value=h);x.font=Font(name='Arial',size=11,bold=True,color='FFFFFF');x.fill=PatternFill('solid',fgColor=NAVY);x.alignment=Alignment(wrap_text=True,vertical='center');x.border=bd
    c.column_dimensions[chr(64+j)].width=15
c.row_dimensions[2].height=45
ramp=[2,2,3,3,4,4,5,5,6,6,6,6,6]
for w in range(1,14):
    r=w+2
    vals=[w,f'=Inputs!$B$6+7*(A{r}-1)',ramp[w-1],f'=C{r}*Inputs!$B$10',f'=D{r}*Inputs!$B$9',
      f'=IF(A{r}-Inputs!$B$24>=1,INDEX($E$3:$E$15,A{r}-Inputs!$B$24),0)',
      f'=D{r}*Inputs!$B$8*(1+Inputs!$B$11+Inputs!$B$12+Inputs!$B$13)',f'=D{r}*Inputs!$B$14',
      (f'=G{r-1}' if w>1 else 0),'=Inputs!$B$15',f'=H{r}+I{r}+J{r}',
      (f'=L{r-1}+F{r}-K{r}' if w>1 else f'=Inputs!$B$7+F{r}-K{r}')]
    for j,v in enumerate(vals,1):
        x=c.cell(row=r,column=j,value=v);x.font=A();x.border=bd
        x.number_format='DD/MM/YYYY' if j==2 else ('0' if j in (1,3,4) else '£#,##0')
    c.cell(row=r,column=3).fill=PatternFill('solid',fgColor=INP);c.cell(row=r,column=3).font=A(color='1F3F99')
tr=16;c.cell(row=tr,column=1,value='Total').font=A(bold=True)
for j in range(4,12):
    col=chr(64+j);x=c.cell(row=tr,column=j,value=f'=SUM({col}3:{col}15)');x.font=A(bold=True);x.number_format='£#,##0' if j>4 else '0';x.border=bd
c.conditional_formatting.add('L3:L15',CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor='F8D7D3'),font=Font(name='Arial',color='9C1C1C',bold=True)))
c.freeze_panes='C3'
c['A18']='Wages for the final week (week 13) are paid in week 14, so they do not appear here. Workers must be paid even if a client pays late or not at all (Conduct Regulations, regulation 12).';c['A18'].font=A(color='555555')
wb.save(OUT)
