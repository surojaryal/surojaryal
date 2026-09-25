import json,csv,datetime
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
OUT="Haverton Recruitment/8 First Client/"
W=json.load(open('week.json'))
MON=datetime.date(2026,9,28)
def due(i,o):
    if o['tier']=='A': return MON
    return MON+datetime.timedelta(days=1 if i<12 else 2)
NAVY='1B2B45';INP='FFF6D9'
A=lambda **k:Font(name='Arial',size=11,**k)
thin=Side(style='thin',color='D9D2C3');bd=Border(top=thin,bottom=thin,left=thin,right=thin)
wb=Workbook();s=wb.active;s.title='This Week'
s['A1']='This Week: 20 Providers to Contact';s['A1'].font=Font(name='Arial',size=16,bold=True,color=NAVY)
s['A2']='Week starting Monday 28 September 2026 | Priority A = live vacancy seen on Indeed, call first | Check every number on the TPS website (Corporate TPS) before calling';s['A2'].font=A(color='555555')
H=['#','Priority','Call on','Service','Provider (legal entity)','Type','Town','Postcode','Miles from Swanley','Phone','Email found','Website','CQC page','Why this one','CTPS checked?','Call 1 outcome','Email sent (date)','Follow up on','Manager name','Notes']
Wd=[4,8,11,34,34,18,13,10,9,14,28,30,40,50,12,20,14,12,36,30]
for j,h in enumerate(H,1):
    c=s.cell(row=4,column=j,value=h);c.font=Font(name='Arial',size=11,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=NAVY if j<=14 else '8A6D2F');c.alignment=Alignment(wrap_text=True,vertical='center');c.border=bd
    s.column_dimensions[c.column_letter].width=Wd[j-1]
for i,o in enumerate(W):
    r=5+i
    vals=[i+1,o['tier'],due(i,o),o['name'],o['provider'],o['type'],o['town'],o['pc'],o['miles'],o['phone'] or 'Not listed: use website',o['email'] or 'Ask on the call',o['web'],o['url'],o['why']+('. No Registered Manager listed on CQC: lead with manager recruitment' if o.get('norm') else ''),'','','',None,(o.get('contact','')+(' (Registered Manager)' if o.get('rm') and o.get('contact')==o.get('rm') else ' (Nominated Individual)' if o.get('contact') else '')),'']
    for j,v in enumerate(vals,1):
        c=s.cell(row=r,column=j,value=v);c.font=A();c.border=bd;c.alignment=Alignment(vertical='top',wrap_text=j in (4,5,14))
        if j==3:c.number_format='ddd DD/MM'
        if j==13 and v:c.hyperlink=v;c.font=A(color='1F4E9A',underline='single')
        if j>=15:c.fill=PatternFill('solid',fgColor=INP)
    s.cell(row=r,column=18,value=f'=IF(Q{r}="","",Q{r}+7)').number_format='DD/MM/YYYY'
    s.cell(row=r,column=17).number_format='DD/MM/YYYY'
    if o['tier']=='A':
        for j in (1,2,4): s.cell(row=r,column=j).font=A(bold=True)
dv1=DataValidation(type='list',formula1='"Yes: not on CTPS,On CTPS: email only"',allow_blank=True);s.add_data_validation(dv1);dv1.add('O5:O24')
dv2=DataValidation(type='list',formula1='"No answer,Left message,Spoke: send info,Meeting booked,Has a vacancy,Not now: call in 60 days,Not interested,Do not contact"',allow_blank=True);s.add_data_validation(dv2);dv2.add('P5:P24')
s.freeze_panes='E5';s.auto_filter.ref='A4:T24';s.row_dimensions[4].height=34
r=26
s.cell(row=r,column=1,value='Weekly targets').font=Font(name='Arial',size=13,bold=True,color=NAVY)
tg=[('Calls made','=COUNTA(P5:P24)',20),('Emails sent','=COUNT(Q5:Q24)',20),('Meetings booked','=COUNTIF(P5:P24,"Meeting booked")',3),('Vacancies found','=COUNTIF(P5:P24,"Has a vacancy")',1)]
s.cell(row=r+1,column=4,value='Measure').font=A(bold=True);s.cell(row=r+1,column=5,value='Done').font=A(bold=True);s.cell(row=r+1,column=6,value='Target').font=A(bold=True)
for k,(a,f,t) in enumerate(tg):
    s.cell(row=r+2+k,column=4,value=a).font=A();c=s.cell(row=r+2+k,column=5,value=f);c.font=A(bold=True);s.cell(row=r+2+k,column=6,value=t).font=A()
p=wb.create_sheet('Week Plan')
plan=[('Day','What to do','Time'),
('Mon 28 Sep','Check all 20 numbers on the TPS website. Call the 5 Priority A providers (live vacancies). Send the matching email from "This Week Emails" straight after each call, or if nobody answers.','3 hours'),
('Tue 29 Sep','Call providers 6 to 12. Send emails after each call.','2.5 hours'),
('Wed 30 Sep','Call providers 13 to 20. Send emails after each call.','2.5 hours'),
('Thu 1 Oct','Discovery meetings or calls that got booked. Send Client Terms and Fee Proposal Letter within 24 hours of any interested provider. Post Advert 1 on Facebook and LinkedIn.','2 hours'),
('Fri 2 Oct','Weekly review: update CRM statuses, fill the targets box, pick next week\'s 20 from Target Providers.','1 hour'),
('Mon 5 Oct','Send follow-up email to everyone who has not replied (Email 3 in the pack). Second call to anyone who said "call back".','2 hours')]
for i,row in enumerate(plan,1):
    for j,v in enumerate(row,1):
        c=p.cell(row=i,column=j,value=v);c.font=A(bold=(i==1),color='FFFFFF' if i==1 else '333333');c.border=bd;c.alignment=Alignment(wrap_text=True,vertical='top')
        if i==1:c.fill=PatternFill('solid',fgColor=NAVY)
p.column_dimensions['A'].width=14;p.column_dimensions['B'].width=100;p.column_dimensions['C'].width=12
wb.save(OUT+'This Week 20 Providers.xlsx')
with open(OUT+'Follow Ups Import.csv','w',newline='',encoding='utf-8') as f:
    w=csv.writer(f);w.writerow(['Type','Organisation / Candidate','Contact','Action','Due Date','Owner','Status','Next Action'])
    for i,o in enumerate(W):
        act=('Call: '+o['why']+'. ' if o['tier']=='A' else 'First call. ')+'Check CTPS first. Then send email '+str(i+1)+' from This Week Emails.'
        w.writerow(['Client',o['name'],o['phone'] or o['web'],act,due(i,o).isoformat(),'Suroj Aryal','Open','Book a 20-minute discovery call'])
    w.writerow(['Internal','Haverton Recruitment','','Weekly review: CRM statuses, targets, pick next 20','2026-10-02','Suroj Aryal','Open',''])
    w.writerow(['Internal','Haverton Recruitment','','Follow-up email to all providers who have not replied (Email 3)','2026-10-05','Suroj Aryal','Open',''])
print('ok')
