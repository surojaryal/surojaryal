import csv,re,datetime
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
OUT="Haverton Recruitment/7 Launch Kit/"
rows=list(csv.reader(open('cqc.csv',encoding='utf-8-sig')))[5:]
want=['Nursing homes','Residential homes','Homecare agencies','Supported living']
core=re.compile(r'^(BR[1-8]|DA\d{1,2}|TN(9|1[0-5]))\s')
def stype(s):
    t=s.split('|')
    if 'Nursing homes' in t: return 'Nursing Home'
    if 'Residential homes' in t: return 'Residential Care'
    if 'Homecare agencies' in t: return 'Domiciliary Care'
    return 'Supported Living'
def phone(p):
    p=p.strip()
    return ('0'+p) if p and p.isdigit() and not p.startswith('0') else p
def date(d):
    try: return datetime.datetime.strptime(d.split(' - ')[0],'%d/%b/%Y').date()
    except: return None
sel=[r for r in rows if r[10] in ('Kent','Bexley','Bromley','Medway') and set(want)&set(r[6].split('|'))]
order={'Nursing Home':0,'Residential Care':1,'Domiciliary Care':2,'Supported Living':3}
def rec(r):
    return dict(area='Priority area' if core.match(r[3]) else 'Wider Kent and Medway',name=r[0],provider=r[9],type=stype(r[6]),types=r[6].replace('|',', '),
      spec=r[8].replace('|',', '),addr=r[2].replace(',',', ').replace(',  ',', '),pc=r[3],la=r[10],phone=phone(r[4]),web=r[5],locid=r[13],url=r[12],checked=date(r[7]))
recs=sorted((rec(r) for r in sel),key=lambda x:(x['area']!='Priority area',x['pc'].split()[0],order[x['type']],x['name']))
pri=[x for x in recs if x['area']=='Priority area']; wid=[x for x in recs if x['area']!='Priority area']
NAVY='1B2B45';GOLD='C9A55C';LIGHT='F3EFE6'
hf=Font(name='Arial',bold=True,color='FFFFFF',size=11);bf=Font(name='Arial',size=11)
thin=Side(style='thin',color='D9D2C3')
wb=Workbook();ws=wb.active;ws.title='How To Use'
guide=[('Target Providers',True),('Haverton Recruitment And Staffing | Haverton Care Limited',False),('',False),
('What this is',True),
(f'{len(recs)} CQC-registered care homes, nursing homes, home care agencies and supported living services in Kent, Medway, Bexley and Bromley.',False),
('Source: CQC care directory published 23 September 2026 (cqc.org.uk, Using CQC data). Open Government Licence. Business contact details only; no named individuals.',False),
(f'"Priority Area" sheet: {len(pri)} services in BR, DA and TN9 to TN15 postcodes (Swanley, Dartford, Bexley, Bromley, Orpington, Sidcup, Gravesend, Sevenoaks, Tonbridge).',False),
(f'"Wider Kent And Medway" sheet: {len(wid)} services further out, for later.',False),('',False),
('How to use it',True),
('1. Filter by postcode or service type. Start with nursing and residential homes near Swanley: they use the most agency and permanent staff.',False),
('2. Before contacting, open the CQC page link to check the service is active and read the latest report.',False),
('3. Find the right contact (Registered Manager or owner) from the service website or by phone. Record only their business contact details.',False),
('4. Log each contact in the Contacted, Date, Outcome and Next action columns, or import into Haverton Operations (Clients, Import CSV) using "Target Providers Import.csv".',False),('',False),
('Marketing rules (legal requirements)',True),
('Calls: before a marketing call, check the number is not on the Corporate Telephone Preference Service (CTPS). Do not call CTPS numbers unless the business has told you it is happy to be called. (PECR)',False),
('Emails: you can email a limited company or other corporate body without prior consent, but every email must say who you are and give an easy way to opt out. Sole traders and partnerships need consent first. Keep an opt-out list. (PECR)',False),
('Tell each organisation where you got their details when you first contact them, and remove anyone who objects. (UK GDPR)',False),
('Do not use the CQC rating or report to criticise a provider in your approach. Lead with how you can help.',False)]
for i,(t,b) in enumerate(guide,1):
    c=ws.cell(row=i,column=1,value=t);c.font=Font(name='Arial',size=16 if i==1 else 11,bold=b,color=NAVY if b else '333333');c.alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=120
H=['Priority area?','Service name','Provider (legal entity)','Main type','All service types','Specialisms','Address','Postcode','Local authority','Phone','Website','CQC location ID','CQC page','Last CQC check','Contacted?','Date contacted','Outcome','Next action']
W=[16,34,34,17,30,40,44,11,15,15,30,15,40,14,13,14,28,28]
keys=['area','name','provider','type','types','spec','addr','pc','la','phone','web','locid','url','checked']
for title,data in (('Priority Area',pri),('Wider Kent And Medway',wid)):
    s=wb.create_sheet(title)
    for j,h in enumerate(H,1):
        c=s.cell(row=1,column=j,value=h);c.font=hf;c.fill=PatternFill('solid',fgColor=NAVY if j<=14 else '8A6D2F');c.alignment=Alignment(wrap_text=True,vertical='center')
        s.column_dimensions[get_column_letter(j)].width=W[j-1]
    for i,x in enumerate(data,2):
        for j,k in enumerate(keys,1):
            c=s.cell(row=i,column=j,value=x[k]);c.font=bf
            if k=='url' and x[k]: c.hyperlink=x[k];c.font=Font(name='Arial',size=11,color='1F4E9A',underline='single')
            if k=='checked': c.number_format='DD/MM/YYYY'
        for j in (16,): s.cell(row=i,column=j).number_format='DD/MM/YYYY'
    dv=DataValidation(type='list',formula1='"Not yet,Emailed,Called,Meeting booked,Terms sent,Client,Not interested,Do not contact"',allow_blank=True)
    s.add_data_validation(dv);dv.add(f'O2:O{len(data)+1}')
    s.freeze_panes='C2';s.auto_filter.ref=f'A1:R{len(data)+1}';s.row_dimensions[1].height=32
wb.save(OUT+'Target Providers.xlsx')
# CRM import: Clients register headings
ch=['Legal Entity','Trading Name / Service','Service Type','CQC Regulated?','CQC Location ID / Note','Phone','Postcode','Client Status','Conflict Of Interest Note']
with open(OUT+'Target Providers Import.csv','w',newline='',encoding='utf-8') as f:
    w=csv.writer(f);w.writerow(ch)
    for x in pri: w.writerow([x['provider'],x['name'],x['type'],'Yes',f"{x['locid']} {x['url']}",x['phone'],x['pc'],'Lead',''])
cand=['Full Name','Email','Phone','Postcode','Target Role','Candidate Route','Current Stage','Source Of Data','Privacy Notice Given','Privacy Notice Version / Date','Consent To Represent','Consent Date','Consent Scope','Availability','Marketing / Contact Preference','Last Contact','Next Action Date']
with open(OUT+'Candidates Import Template.csv','w',newline='',encoding='utf-8') as f: csv.writer(f).writerow(cand)
import collections
print(len(recs),len(pri),len(wid),collections.Counter(x['type'] for x in pri))
